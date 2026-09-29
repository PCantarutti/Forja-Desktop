import asyncio
import io
import json
import re

import pytest

from app import db, design, design_export, design_html, mirror
from app.tools import ToolError

M = {k: {"provider": "ollama", "model": "m"} for k in ("plano", "geracao", "edicao")}
DOC = """<!doctype html><html><head><style>
:root { --cor-primaria: #c75b12; --cor-fundo: #ffffff; --fonte-titulo: Georgia, serif; }
</style></head><body><section data-section="hero"><h1>Padaria</h1></section></body></html>"""


@pytest.fixture(autouse=True)
def pastas(tmp_path, monkeypatch):
    monkeypatch.setattr(mirror, "ROOT", tmp_path / "conversas")
    design._RUNS.clear()
    yield
    design._RUNS.clear()


def _projeto(html=DOC) -> int:
    with db.session() as s:
        c = db.Conversation(kind="design", title="Padaria")
        s.add(c)
        s.commit()
        cid = c.id
    design._nova_versao(cid, None, html, "inicial")
    return cid


def _esperar(fn):
    async def main():
        r = fn()
        await asyncio.gather(*[t for t in asyncio.all_tasks() if t is not asyncio.current_task()])
        return r
    return asyncio.run(main())


def test_variacoes_so_com_tokens_existentes_e_escolha_sem_ia(monkeypatch):
    conv = _projeto()
    vistos = []

    async def chat_stream(provider, model, messages, tools, num_ctx, effort=None, **kw):
        vistos.append(messages)
        yield "content", json.dumps({"mensagem": "Três caminhos.", "variacoes": [
            {"nome": "Oceano", "descricao": "frio", "tokens": {"--cor-primaria": "#0b4f6c", "--cor-inventada": "#fff"}},
            {"nome": "Noite", "tokens": {"--cor-fundo": "#111111", "--cor-primaria": "#f2aa4c"}},
            {"nome": "Vazia", "tokens": {"--nao-existe": "red"}},
            {"nome": "Ruim", "tokens": {"--cor-primaria": "red; } body{display:none"}}]})
        yield "done", {}
    monkeypatch.setattr(design.llm, "chat_stream", chat_stream)
    msg = _esperar(lambda: design.start(conv, "", M, rota="variacoes"))
    assert re.match(r"<!-- variacoes", vistos[0][0]["content"]) and "--cor-primaria: #c75b12" in vistos[0][1]["content"]
    m = design.projeto(conv)["mensagens"][-1]
    assert m["status"] == "variacoes" and [v["nome"] for v in m["variacoes"]] == ["Oceano", "Noite"]
    assert m["variacoes"][0]["tokens"] == {"--cor-primaria": "#0b4f6c"} and design.projeto(conv)["total"] == 1

    monkeypatch.setattr(design.llm, "chat_stream", lambda *a, **k: (_ for _ in ()).throw(AssertionError("chamou o modelo")))
    r = design.escolher_variacao(conv, msg["id"], 1)
    p = r["projeto"]
    assert p["atual"] == 2 and "--cor-fundo: #111111" in p["html"] and "--cor-primaria: #f2aa4c" in p["html"]
    assert p["mensagens"][-1]["content"] == "v2: variação: Noite" and "Noite" in p["mensagens"][-1]["passos"][0]
    assert [x for x in p["mensagens"] if x["id"] == msg["id"]][0]["escolhida"] == 1
    with pytest.raises(ToolError):
        design.escolher_variacao(conv, msg["id"], 9)


def test_html_de_qualquer_versao():
    conv = _projeto()
    design.ajustar_tokens(conv, {"--cor-primaria": "#000000"})
    v1, v2 = design.versao_html(conv, 1), design.versao_html(conv, 2)
    assert "#c75b12" in v1["html"] and "#000000" in v2["html"] and v1["descricao"] == "inicial"
    with pytest.raises(ToolError):
        design.versao_html(conv, 7)


def _chromium() -> bool:
    async def abre():
        from playwright.async_api import async_playwright
        async with async_playwright() as pw:
            await (await pw.chromium.launch(headless=True)).close()
    try:
        asyncio.run(abre())
        return True
    except Exception:
        return False


@pytest.mark.skipif(not _chromium(), reason="Chromium do Playwright não instalado")
def test_pptx_um_slide_por_slide_com_notas():
    from pptx import Presentation
    plano = {"tipo": "slides", "titulo": "Pitch", "tokens": {"--cor-fundo": "#fff"},
             "secoes": [{"nome": f"s{i}"} for i in range(1, 4)]}
    h = design_html.esqueleto(plano)
    for i in range(1, 4):
        h = h.replace(f'data-section="s{i}" data-slide data-placeholder="1">Gerando o slide “s{i}”…',
                      f'data-section="s{i}" data-slide><h2>Slide número {i}</h2>')
    dados, tipo, nome = asyncio.run(design_export.exportar(h, "Pitch", "pptx"))
    prs = Presentation(io.BytesIO(dados))
    assert nome == "pitch.pptx" and "presentation" in tipo and len(prs.slides) == 3
    assert round(prs.slide_width / prs.slide_height, 2) == round(16 / 9, 2)
    assert prs.slides[1].notes_slide.notes_text_frame.text == "Slide número 2"
    assert prs.slides[0].shapes[0].shape_type == 13   # PICTURE
    with pytest.raises(ToolError):
        asyncio.run(design_export.exportar(DOC, "Site", "pptx"))
